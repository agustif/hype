#!/usr/bin/env python3
"""Package Hype's rendered slides as a self-contained PowerPoint file."""
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile

from pptx import Presentation
from pptx.util import Inches


def export(manifest, destination):
    manifest = Path(manifest)
    deck = json.loads(manifest.read_text())
    result = Presentation()
    result.slide_width = Inches(13.333333)
    result.slide_height = Inches(7.5)
    result.core_properties.title = deck['title']
    result.core_properties.author = 'Hype'
    scale = result.slide_width / 1920
    for entry in deck['slides']:
        slide = result.slides.add_slide(result.slide_layouts[6])
        slide.shapes.add_picture(str(manifest.parent / entry['image']), 0, 0,
                                 result.slide_width, result.slide_height)
        if not entry.get('video'):
            continue
        movie = Path(entry['video'])
        probe = json.loads(subprocess.check_output([
            'ffprobe', '-v', 'error', '-show_streams', '-of', 'json', str(movie)], text=True))
        video = next((s for s in probe['streams'] if s['codec_type'] == 'video'), None)
        audio = next((s for s in probe['streams'] if s['codec_type'] == 'audio'), None)
        if movie.suffix.lower() != '.mp4' or not video or video['codec_name'] != 'h264' or (audio and audio['codec_name'] != 'aac'):
            raise ValueError(f'PowerPoint video requires H.264/AAC MP4: {movie.name}')
        if entry['span']:
            x, y, w, h = 0, 0, 1920, 1080
        elif entry['title']:
            x, y, w, h = 100, 280, 1720, 730
        else:
            x, y, w, h = 70, 50, 1780, 980
        ratio = video['width'] / video['height']
        if not entry['span']:
            width, height = min(w, h*ratio), min(h, w/ratio)
            x += (w-width)/2; y += (h-height)/2; w,h = width,height
        elif abs(ratio - 16/9) > .01:
            raise ValueError(f'Spanning video must be 16:9 for this first PowerPoint exporter: {movie.name}. Use fit.')
        shape = slide.shapes.add_movie(str(movie), int(x*scale), int(y*scale),
                                      int(w*scale), int(h*scale),
                                      poster_frame_image=entry['poster'], mime_type='video/mp4')
        # python-pptx creates the media timing node. Set its existing playback
        # controls without inventing a separate animation tree.
        ns = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main'}
        nodes = slide._element.findall('.//p:video/p:cMediaNode', ns)
        for node in nodes:
            target = node.find('./p:tgtEl/p:spTgt', ns)
            if target is None or target.get('spid') != str(shape.shape_id):
                continue
            node.set('vol', '0' if entry['muted'] else '100000')
            timing = node.find('{http://schemas.openxmlformats.org/presentationml/2006/main}cTn')
            if entry['loop']: timing.set('repeatCount', 'indefinite')
            for condition in timing.findall('./p:stCondLst/p:cond', ns):
                condition.set('delay', '0' if entry['autoplay'] else 'indefinite')
        if entry.get('overlay_image'):
            slide.shapes.add_picture(str(manifest.parent/entry['overlay_image']), 0, 0,
                                     result.slide_width, result.slide_height)
    destination = Path(destination).resolve()
    destination.parent.mkdir(parents=True, exist_ok=True)
    fd, temporary = tempfile.mkstemp(prefix='.hype-', suffix='.pptx', dir=destination.parent)
    os.close(fd)
    try:
        result.save(temporary)
        os.replace(temporary, destination)
    finally:
        if os.path.exists(temporary): os.unlink(temporary)

if __name__ == '__main__':
    try:
        export(sys.argv[1], sys.argv[2])
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
